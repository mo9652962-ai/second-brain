# -*- coding: utf-8 -*-
"""扇叶开场 v3.0（三维透视版）演示文稿生成器

【深挖阶段2 突破】原版视频真正的"精致感"来源：
PowerPoint「三维旋转（透视：前部）」——扇叶在 3D 空间倾斜，产生真实空间纵深。
本生成器用数学三维投影（rot_x/rot_z + 透视投影）复现该效果。

对齐原版全部参数：
- 扇叶：6 片缺角圆形，左上重合→逐片旋转 60°→全选组合
- 三维旋转：透视：前部（X≈+18~30°）
- 柔化边缘：10 磅
- 阴影：透明度 15%、大小 104%、模糊 5 磅、距离 0、角度 0
- 填充：幻灯片背景填充（扇叶透出背景）
- 背景：绿色系庭院（#9FBD8A / #6C8F57 / #334A2B）
- 收尾：平滑切换（Morph）
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml import parse_xml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import render_fan_3d as f3


def add_morph_transition(slide):
    xml = (
        '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'spd="med" advClick="1"><p:morph option="byObject"/></p:transition>'
    )
    slide._element.append(parse_xml(xml))


def build_3d_presentation(tilt_x=22, tilt_z=0, out_name="Fan_3D_Opening.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    tmp = os.path.join(os.environ.get("LOCALAPPDATA", "C:/Temp"), "Temp")
    os.makedirs(tmp, exist_ok=True)

    bg = f3.create_courtyard_bg_v3(1920, 1080)

    # Slide 1：合拢态
    s1_img = f3.render_fan_3d_v2(bg, False, tilt_x=tilt_x, tilt_z=tilt_z)
    p1 = os.path.join(tmp, "fan3d_s1.jpg")
    s1_img.save(p1, quality=95)
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(p1, 0, 0, width=Inches(13.333), height=Inches(7.5))

    # Slide 2：展开态 + Morph
    s2_img = f3.render_fan_3d_v2(bg, True, tilt_x=tilt_x, tilt_z=tilt_z)
    p2 = os.path.join(tmp, "fan3d_s2.jpg")
    s2_img.save(p2, quality=95)
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_picture(p2, 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_morph_transition(s2)

    out_dir = os.path.join(os.path.expanduser("~"), ".openclaw", "workspace", "knowledge", "Productivity")
    out = os.path.join(out_dir, out_name)
    prs.save(out)
    print(f"✅ 三维透视扇叶开场 PPT 生成: {out}")
    return out


if __name__ == "__main__":
    build_3d_presentation()
