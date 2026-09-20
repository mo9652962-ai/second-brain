# -*- coding: utf-8 -*-
"""回归测试：render_fan_replica_lib.py（扇叶开场忠实还原渲染库 v2.8+）

覆盖：
1. create_courtyard_background: 尺寸/模式 + 增强层像素抽查（水洼反光/轻烟带）
2. render_fan_replica: 双态渲染 + 状态差 + 细节层存在性（金边菱形/孤舟剪影）
3. _polar: PIL 坐标转换正确性（与扇叶 pieslice 坐标系一致）
4. 端到端 PPT: Fan_Replica_GuFeng.pptx 结构（2页/16:9/Morph/背景图）

运行：python tests/test_fan_replica_lib.py   （退出码 0=通过）
"""
import sys, os
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, SCRIPT_DIR)
import render_fan_replica_lib as lib

P = F = 0
def check(name, cond, info=""):
    global P, F
    if cond:
        P += 1
        print(f"  ✅ {name}")
    else:
        F += 1
        print(f"  ❌ {name} {info}")

print("=== 1. 背景生成（增强层）===")
bg = lib.create_courtyard_background(1920, 1080)
check("背景 1920x1080 RGB", isinstance(bg, Image.Image) and bg.size == (1920, 1080) and bg.mode == "RGB")

arr = np.asarray(bg)
puddle_region = arr[int(1080*0.85):1080, int(1920*0.40):int(1920*0.60)]
check("水洼反光区亮度 > 全图均值", puddle_region.mean() > arr.mean() + 3,
      f"puddle={puddle_region.mean():.1f} global={arr.mean():.1f}")
smoke_region = arr[int(1080*0.16):int(1080*0.31), int(1920*0.55):int(1920*0.70)]
check("轻烟带存在（偏亮）", smoke_region.mean() > arr.mean() - 2,
      f"smoke={smoke_region.mean():.1f} global={arr.mean():.1f}")

print("\n=== 2. _polar 坐标一致性 ===")
x90, y90 = lib._polar(960, 864, 670, 90)
check("90°(PIL)=正下方（y 增大，与 pieslice 顺时针一致）", y90 > 864 and abs(x90 - 960) < 2, f"({x90:.0f},{y90:.0f})")
x270, y270 = lib._polar(960, 864, 670, 270)
check("270°(PIL)=正上方（y 减小）", y270 < 864 and abs(x270 - 960) < 2, f"({x270:.0f},{y270:.0f})")

print("\n=== 3. 双态渲染 + 细节层存在性 ===")
c = lib.render_fan_replica(bg, is_opened=False)
o = lib.render_fan_replica(bg, is_opened=True)
check("双态尺寸/模式正确", c.size == o.size == (1920, 1080) and o.mode == "RGB")
check("双态状态差有效", c.tobytes() != o.tobytes())

o_arr = np.asarray(o)
cx, cy, radius = 960, 864, 670
gold_count = 0
for deg in range(190, 350, 1):
    for rr_off in range(4, 26):
        bx, by = lib._polar(cx, cy, radius + rr_off, deg)
        x, y = int(bx), int(by)
        if 0 <= x < 1920 and 0 <= y < 1080:
            r, g, b = o_arr[y, x]
            if r > 140 and g > 105 and b < 140:
                gold_count += 1
check("扇缘金边菱形存在（金色像素≥300）", gold_count >= 300, f"gold={gold_count}")
boat_region = o_arr[int(1080*0.62):int(1080*0.80), int(1920*0.82):int(1920*0.97)]
dark_px = ((boat_region[:, :, 0] < 95) & (boat_region[:, :, 1] < 105)).sum()
check("孤舟剪影存在（右下深色像素>200）", dark_px > 200, f"dark={dark_px}")

print("\n=== 4. 端到端 PPT 结构 ===")
pptx_path = os.path.join(os.path.dirname(SCRIPT_DIR), "Fan_Replica_GuFeng.pptx")
check("成品存在且体积有效", os.path.exists(pptx_path) and os.path.getsize(pptx_path) > 300000)
prs = Presentation(pptx_path)
check("2 页（合拢→展开）", len(prs.slides) == 2)
check("16:9 宽屏", round(prs.slide_width.inches, 2) == 13.33 and round(prs.slide_height.inches, 2) == 7.5)
s2 = prs.slides[1]
trans = s2._element.find(qn('p:transition'))
morph = trans.find(qn('p:morph')) if trans is not None else None
check("Slide 2 p:morph 平滑引擎", morph is not None)
for i, s in enumerate(prs.slides):
    n_pic = len(s.shapes._spTree.findall(qn('p:pic')))
    check(f"Slide {i+1} 含背景图", n_pic >= 1, f"got {n_pic}")

print(f"\n{'='*40}\n通过 {P} / 失败 {F}")
sys.exit(1 if F else 0)
